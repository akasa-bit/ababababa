import openpyxl, sys, math, statistics as st
wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
info = list(wb["BCL v2.7"].iter_rows(values_only=True))
for r in info[:26]:
    t=[c for c in r if c]; 
    if t: print("INFO", t)
def load(name):
    rows = list(wb[name].iter_rows(values_only=True)); hdr = rows[0]
    d = {}
    for r in rows[1:]:
        if r[0] is None: continue
        try: y = int(str(r[0])[:4])
        except: continue
        for c, v in zip(hdr[1:], r[1:]):
            if c and isinstance(v, (int, float)): d.setdefault(c, {})[y] = v
    return d
T = load("TFP"); LP = load("Labor Productivity")
print("TFP countries", sorted(T), "years", min(min(v) for v in T.values()), max(max(v) for v in T.values()))
eras = [(1890,1913),(1913,1950),(1950,1974),(1974,1995),(1995,2005),(2005,2019),(2019,2024),(1890,2024)]
def g(d,c,a,b):
    try: return 100*math.log(d[c][b]/d[c][a])/(b-a)
    except (KeyError, ValueError, ZeroDivisionError): return None
print("\nTFP growth %/yr (log) by era")
for c in ["USA","GBR","FRA","DEU","JPN","ITA","ESP","CAN","SWE","KOR","AUS","CHL","MEX","ZAF"] + ["EURO","Euro Area","EA"]:
    if c in T: print(f"  {c}: " + " ".join(f"{a}-{b}:{g(T,c,a,b):.2f}" if g(T,c,a,b) is not None else f"{a}-{b}:NA" for a,b in eras))
print("\nMedian across countries (TFP | LP per hour)")
for a,b in eras:
    v=[g(T,c,a,b) for c in T]; v=[x for x in v if x is not None]
    w=[g(LP,c,a,b) for c in LP]; w=[x for x in w if x is not None]
    print(f"  {a}-{b}: TFP median {st.median(v):.2f} (n={len(v)}) | LP median {st.median(w):.2f} (n={len(w)})")
# TFP level relative to US over time (level index comparability: 'TFP' in PPP? check)
print("\nTFP level ratio to USA (if levels comparable):")
for y in [1890,1913,1950,1974,1995,2019,2024]:
    r=[T[c][y]/T["USA"][y] for c in T if c!="USA" and y in T[c] and y in T["USA"]]
    if r: print(f"  {y}: median {st.median(r):.2f} min {min(r):.2f} max {max(r):.2f} n={len(r)}")
